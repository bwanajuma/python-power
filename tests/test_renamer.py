import tempfile
import unittest
from pathlib import Path
from typing import cast
from unittest.mock import patch

from powerbills.renamer import (
    extract_account_number,
    extract_bill_data,
    load_station_lookup,
    merge_bill_data,
    process_files,
)


class ExtractAccountNumberTests(unittest.TestCase):
    def test_extracts_labelled_number_with_spacing(self):
        self.assertEqual(extract_account_number("Account Number: 123 456 789"), "123456789")

    def test_uses_account_row_when_ocr_drops_punctuation(self):
        self.assertEqual(extract_account_number("ACCOUNT 123456789"), "123456789")

    def test_ignores_unrelated_numbers(self):
        self.assertIsNone(extract_account_number("Invoice 123456789\nMeter 456789"))

    def test_extracts_requested_bill_fields(self):
        bill = extract_bill_data(
            "ACCOUNT NUMBER: 175565209\n"
            "Date Due: 09/09/2026\n"
            "Date of Issue: 23/08/2026\n"
            "Consumption Period: 23/07/2026-22/08/2026\n"
            "ELECTRICITY BILL: Ksh 103.00\n"
            "Fuel Energy Cost: Ksh 1,234.50\n"
            "HighRateConsumption 3kWh x 12.28\n"
            "LowRateConsumption 2kWh x 12.28"
        )
        assert bill is not None
        self.assertEqual(bill.account_number, "175565209")
        self.assertEqual(bill.consumption, 1234.5)
        self.assertEqual(bill.bill, 103.0)
        self.assertEqual(bill.due_date, "09/09/2026")
        self.assertEqual(bill.issue_date, "23/08/2026")
        self.assertEqual(bill.low_rate_consumption, 2.0)
        self.assertEqual(bill.high_rate_consumption, 3.0)

    def test_extracts_issue_date_label_variant(self):
        bill = extract_bill_data("ACCOUNT NUMBER: 175565209\nIssue Date: 23/08/2026")
        assert bill is not None
        self.assertEqual(bill.issue_date, "23/08/2026")

    def test_extracts_due_date_and_method_of_charge(self):
        bill = extract_bill_data(
            "ACCOUNT NUMBER: 175565209\n"
            "Due Date: 07/10/2026\n"
            "Method of Charge: Small Commercial (SC3) "
            "Forex Exchange Adj. (FERFA) 897kWh x 1.1443"
        )
        assert bill is not None
        self.assertEqual(bill.due_date, "07/10/2026")
        self.assertEqual(bill.method_of_charge, "Small Commercial (SC3)")

    def test_ignores_method_of_charge_without_classified_tariff(self):
        bill = extract_bill_data(
            "ACCOUNT NUMBER: 175565209\n"
            "Method of Charge: Forex Exchange Adj. (FERFA) 897kWh x 1.1443"
        )
        assert bill is not None
        self.assertIsNone(bill.method_of_charge)

    def test_extracts_supported_tariff_classes_only(self):
        tariffs = (
            "Domestic Lifeline (0–30 units)",
            "Domestic Ordinary 1 (31–100 units)",
            "Domestic Ordinary 2 (Over 100 units)",
            "Electric Vehicles (E-Mobility)",
        )
        for tariff in tariffs:
            with self.subTest(tariff=tariff):
                bill = extract_bill_data(
                    "ACCOUNT NUMBER: 175565209\n"
                    f"Method of Charge: {tariff} Forex Exchange Adj. (FERFA)"
                )
                assert bill is not None
                self.assertEqual(bill.method_of_charge, tariff)


class ProcessFilesTests(unittest.TestCase):
    def test_preview_does_not_rename(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "bill.pdf"
            source.write_bytes(b"pdf")
            result = process_files([source], ocr=lambda _: "Account Number: 123456789")[0]
            self.assertEqual(result.status, "planned")
            self.assertTrue(source.exists())

    def test_apply_renames_file(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "bill.pdf"
            source.write_bytes(b"pdf")
            result = process_files([source], apply=True, ocr=lambda _: "Account Number: 123456789")[
                0
            ]
            self.assertEqual(result.status, "renamed")
            self.assertTrue((Path(folder) / "123456789.pdf").exists())

    def test_account_header_wins_over_other_numbers(self):
        text = (
            "ACCOUNT NUMBER: 178761435\nInvoice Number: 260910015872316\nMeter Number: 040016111611"
        )
        self.assertEqual(
            process_files([Path("bill.pdf")], ocr=lambda _: text)[0].message, "178761435"
        )

    def test_uses_station_name_from_lookup(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "178761435.pdf"
            source.write_bytes(b"pdf")
            result = process_files(
                [source],
                apply=True,
                ocr=lambda _: "Account Number: 178761435",
                station_lookup={"178761435": "Athi River"},
            )[0]
            self.assertEqual(result.message, "178761435")
            self.assertTrue((Path(folder) / "Athi River.pdf").exists())

    def test_falls_back_to_account_when_station_is_missing(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "bill.pdf"
            source.write_bytes(b"pdf")
            process_files(
                [source],
                apply=True,
                ocr=lambda _: "Account Number: 999999999",
                station_lookup={"178761435": "Athi River"},
            )
            self.assertTrue((Path(folder) / "999999999.pdf").exists())


class WorkbookLookupTests(unittest.TestCase):
    def test_loads_requested_columns(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as folder:
            workbook_path = Path(folder) / "lookup.xlsx"
            workbook = Workbook()
            from openpyxl.worksheet.worksheet import Worksheet

            sheet = cast(Worksheet, workbook.active)
            sheet.title = "September"
            sheet.append(["Station", "Account Number"])
            sheet.append(["Athi River", 178924769])
            workbook.save(workbook_path)
            self.assertEqual(load_station_lookup(workbook_path), {"178924769": "Athi River"})

    def test_loads_the_only_sheet_when_its_name_changes(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as folder:
            workbook_path = Path(folder) / "lookup.xlsx"
            workbook = Workbook()
            from openpyxl.worksheet.worksheet import Worksheet

            sheet = cast(Worksheet, workbook.active)
            sheet.title = "September"
            sheet.append(["Station", "Account Number"])
            sheet.append(["Athi River", 178924769])
            workbook.save(workbook_path)
            self.assertEqual(load_station_lookup(workbook_path), {"178924769": "Athi River"})

    def test_merges_bill_fields_into_template(self):
        from openpyxl import Workbook, load_workbook

        with tempfile.TemporaryDirectory() as folder:
            template = Path(folder) / "template.xlsx"
            output = Path(folder) / "output.xlsx"
            workbook = Workbook()
            from openpyxl.worksheet.worksheet import Worksheet

            sheet = cast(Worksheet, workbook.active)
            sheet.title = "Sheet1"
            sheet.append(["Station", "Meter Number", "Account Number"])
            sheet.append(["Athi River", "123", 178924769])
            workbook.save(template)
            bill = extract_bill_data(
                "ACCOUNT NUMBER: 178924769\nDate Due: 09/09/2026\n"
                "Issue Date: 23/08/2026\n"
                "Consumption Period: 23/07/2026-22/08/2026\n"
                "ELECTRICITY BILL: Ksh 103.00\n"
                "Fuel Energy Cost 1,234.50\n"
                "Method of Charge: Small Commercial (SC3)"
            )
            assert bill is not None
            merge_bill_data(
                template,
                output,
                [bill],
            )
            result = load_workbook(output, data_only=True)["Sheet1"]
            headers = [cell.value for cell in result[1]]
            values = [cell.value for cell in result[2]]
            expected = {
                "Consumption Period": "23/07/2026-22/08/2026",
                "consumption": 1234.5,
                "Bill": 103,
                "Due Date": "09/09/2026",
                "Issue Date": "23/08/2026",
                "Method of Charge": "Small Commercial (SC3)",
                "LowRateConsumption": None,
                "HighRateConsumption": None,
            }
            for header, expected_value in expected.items():
                self.assertEqual(values[headers.index(header)], expected_value)


class ProcessFilesWithLookupTests(unittest.TestCase):
    def test_uses_station_name_from_lookup_file(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "178761435.pdf"
            source.write_bytes(b"pdf")
            lookup_file = Path(folder) / "lookup.xlsx"
            from openpyxl import Workbook

            workbook = Workbook()
            from openpyxl.worksheet.worksheet import Worksheet

            sheet = cast(Worksheet, workbook.active)
            sheet.title = "Sheet1"
            sheet.append(["Station", "Account Number"])
            sheet.append(["Athi River", 178761435])
            workbook.save(lookup_file)
            result = process_files(
                [source],
                apply=True,
                ocr=lambda _: "Account Number: 178761435",
                station_lookup=load_station_lookup(lookup_file),
            )[0]
            self.assertEqual(result.message, "178761435")
            self.assertTrue((Path(folder) / "Athi River.pdf").exists())


class CliDefaultsTests(unittest.TestCase):
    def test_production_input_directory_is_raw_files(self):
        from powerbills.cli import DEFAULT_INPUT_DIRECTORY

        self.assertEqual(DEFAULT_INPUT_DIRECTORY, "raw_files")

    def test_skips_cleanly_when_no_pdfs_are_found(self):
        from powerbills.cli import main

        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "output.xlsx"
            arguments = ["powerbills", folder, "--output", str(output)]
            with patch("sys.argv", arguments):
                self.assertEqual(main(), 0)
            self.assertFalse(output.exists())

    def test_extract_only_writes_fields_without_rename_planning(self):
        from openpyxl import Workbook, load_workbook
        from openpyxl.worksheet.worksheet import Worksheet

        from powerbills.cli import main

        with tempfile.TemporaryDirectory() as folder:
            folder_path = Path(folder)
            source = folder_path / "bill.pdf"
            source.write_bytes(b"pdf")
            template = folder_path / "template.xlsx"
            output = folder_path / "output.xlsx"
            workbook = Workbook()
            sheet = cast(Worksheet, workbook.active)
            sheet.title = "September"
            sheet.append(["Station", "Account Number"])
            sheet.append(["Athi River", "175565209"])
            workbook.save(template)
            ocr_text = (
                "ACCOUNT NUMBER: 175565209\n"
                "Due Date: 07/10/2026\n"
                "Date of Issue: 22/09/2026\n"
                "Consumption Period: 29/08/2026-21/09/2026\n"
                "ELECTRICITY BILL: Ksh 242,441.00\n"
                "Fuel Energy Cost: Ksh 12,345.67\n"
                "LowRateConsumption 5392kWh x 8\n"
                "HighRateConsumption 6605kWh x 16\n"
                "Method of Charge: E-Mobility Levies and Adjustments"
            )
            arguments = [
                "powerbills",
                str(source),
                "--workbook",
                str(template),
                "--output",
                str(output),
                "--extract-only",
            ]
            with (
                patch("sys.argv", arguments),
                patch("powerbills.cli.ocr_pdf", return_value=ocr_text),
                patch(
                    "powerbills.cli.process_files",
                    side_effect=AssertionError("rename processing should be skipped"),
                ),
            ):
                self.assertEqual(main(), 0)

            result = load_workbook(output, data_only=True)["September"]
            headers = [cell.value for cell in result[1]]
            values = [cell.value for cell in result[2]]
            expected = {
                "Bill": 242441,
                "consumption": 12345.67,
                "Consumption Period": "29/08/2026-21/09/2026",
                "LowRateConsumption": 5392,
                "HighRateConsumption": 6605,
                "Method of Charge": "E-Mobility",
                "Due Date": "07/10/2026",
                "Issue Date": "22/09/2026",
            }
            for header, expected_value in expected.items():
                self.assertEqual(values[headers.index(header)], expected_value)
            self.assertTrue(source.exists())


if __name__ == "__main__":
    unittest.main()
