"""Tests for the streaming swap dump aggregator."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pytest

from powerbills import swapdata
from powerbills.swapdata import (
    SwapDumpError,
    aggregate_swap_dump,
    column_letters,
    combine_reports,
    discover_dumps,
    period_bounds,
    period_from_filename,
    serial_to_datetime,
    summary_rows,
    write_csv,
    write_summary,
)

HEADERS = (
    "Swap DateTime",
    "Country",
    "Station",
    "Customer Name",
    "Incoming Battery Previous SOC",
    "Incoming Battery Current SOC",
    "Net SOC",
    "Billable SOC",
    "Swap Cost",
)


def build_dump(path: Path, rows: list[dict[str, object]], sheet: str = "Swap Data") -> Path:
    """Write a swap dump workbook that mimics the KPLC export layout."""
    from openpyxl import Workbook

    book = Workbook()
    sheet_view = book.active
    sheet_view.title = sheet
    sheet_view.append(list(HEADERS))
    for row in rows:
        sheet_view.append([row.get(header) for header in HEADERS])
    book.save(path)
    return path


def sample_rows() -> list[dict[str, object]]:
    """Rows covering two stations, an out-of-period row, and a blank station."""
    return [
        {
            "Swap DateTime": datetime(2026, 8, 1, 0, 0, 3),
            "Station": "Ahero",
            "Net SOC": 10,
            "Billable SOC": 9,
            "Swap Cost": 150.5,
        },
        {
            "Swap DateTime": datetime(2026, 8, 15, 12, 30),
            "Station": "Alendu",
            "Net SOC": 20.5,
            "Billable SOC": 20,
            "Swap Cost": 300,
        },
        {
            "Swap DateTime": datetime(2026, 8, 31, 23, 59, 59),
            "Station": "ahero ",
            "Net SOC": 5,
            "Billable SOC": 5,
            "Swap Cost": 75,
        },
        {
            "Swap DateTime": datetime(2026, 9, 1),
            "Station": "Ahero",
            "Net SOC": 999,
            "Billable SOC": 999,
            "Swap Cost": 999,
        },
        {
            "Swap DateTime": datetime(2026, 8, 20),
            "Station": None,
            "Net SOC": 4,
            "Billable SOC": 4,
            "Swap Cost": 60,
        },
    ]


def test_column_letters_round_trip() -> None:
    assert column_letters(1) == "A"
    assert column_letters(20) == "T"
    assert column_letters(26) == "Z"
    assert column_letters(27) == "AA"
    with pytest.raises(ValueError):
        column_letters(0)


def test_serial_to_datetime_is_exact_at_midnight() -> None:
    assert serial_to_datetime(46265.9999537037).date().isoformat() == "2026-08-31"
    assert serial_to_datetime(46266).time().isoformat() == "00:00:00"


def test_period_bounds_are_half_open() -> None:
    start, end = period_bounds("2026-08")
    assert start == datetime(2026, 8, 1)
    assert end == datetime(2026, 9, 1)
    assert period_bounds("2026-12")[1] == datetime(2027, 1, 1)
    for invalid in ("2026-13", "August", "26-8"):
        with pytest.raises(ValueError):
            period_bounds(invalid)


def test_period_from_filename() -> None:
    assert period_from_filename(Path("August_Swap_Data.xlsx"), 2026) == "2026-08"
    assert period_from_filename(Path("september_swap.xlsx"), 2026) == "2026-09"
    assert period_from_filename(Path("swap_dump.xlsx"), 2026) is None


def test_aggregate_totals_only_in_period(tmp_path: Path) -> None:
    dump = build_dump(tmp_path / "August_Swap_Data.xlsx", sample_rows())
    report = aggregate_swap_dump(dump, year=2026)

    assert report.sheet == "Swap Data"
    assert report.period == "2026-08"
    assert report.rows_scanned == 5
    assert report.rows_in_period == 4
    assert report.rows_outside_period == 1
    assert set(report.stations) == {"Ahero", "ahero", "Alendu"}
    assert report.swaps == 3
    assert report.net_soc == pytest.approx(35.5)
    assert report.billable_soc == pytest.approx(34.0)
    assert report.swap_cost == pytest.approx(525.5)
    assert report.observed_end == datetime(2026, 9, 1)
    assert report.stations["Ahero"].first_swap == datetime(2026, 8, 1, 0, 0, 3)


def test_aggregate_uses_the_largest_matching_sheet(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    dump = build_dump(
        tmp_path / "September_Swap_Data.xlsx", sample_rows(), sheet="September_Swap_Data"
    )
    book = load_workbook(dump)
    pivot = book.create_sheet("Sheet1")
    pivot.append(["Swap DateTime", "(All)", "Station"])
    pivot.append([None, None, "Row Labels"])
    book.save(dump)
    book.close()

    report = aggregate_swap_dump(dump, year=2026)
    assert report.sheet == "September_Swap_Data"


def test_aggregate_rejects_workbooks_without_swap_headers(tmp_path: Path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "wrong.xlsx"
    book = Workbook()
    book.active.append(["Station", "Notes"])
    book.save(path)
    book.close()

    with pytest.raises(SwapDumpError, match="Net SOC"):
        aggregate_swap_dump(path)


def test_aggregate_handles_cells_split_across_chunks(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    dump = build_dump(tmp_path / "August_Swap_Data.xlsx", sample_rows() * 40)
    whole = aggregate_swap_dump(dump, year=2026)

    monkeypatch.setattr(swapdata, "CHUNK_BYTES", 37)
    chunked = aggregate_swap_dump(dump, year=2026)

    assert chunked.rows_scanned == whole.rows_scanned
    assert chunked.stations.keys() == whole.stations.keys()
    assert chunked.net_soc == pytest.approx(whole.net_soc)
    assert chunked.billable_soc == pytest.approx(whole.billable_soc)
    assert chunked.swap_cost == pytest.approx(whole.swap_cost)


def test_missing_file_raises(tmp_path: Path) -> None:
    with pytest.raises(SwapDumpError, match="not found"):
        aggregate_swap_dump(tmp_path / "nope.xlsx")


def test_summary_rows_and_workbook(tmp_path: Path) -> None:
    dump = build_dump(tmp_path / "August_Swap_Data.xlsx", sample_rows())
    report = aggregate_swap_dump(dump, year=2026)
    rows = summary_rows([report])
    alendu = next(row for row in rows if row["Swap Station"] == "Alendu")

    assert alendu["Net SOC"] == pytest.approx(20.5)
    assert alendu["Billable SOC"] == pytest.approx(20.0)
    assert alendu["Swap Cost"] == pytest.approx(300.0)
    assert alendu["Average Net SOC"] == pytest.approx(20.5)
    assert all(row["Period"] == "2026-08" for row in rows)

    output = write_summary(rows, [report], tmp_path / "out.xlsx")
    assert output.is_file()
    from openpyxl import load_workbook

    book = load_workbook(output, read_only=True)
    assert book.sheetnames == ["Station Summary", "Metadata"]
    assert next(book["Station Summary"].iter_rows(values_only=True)) == (
        "Period",
        "Swap Station",
        "Swaps",
        "Net SOC",
        "Billable SOC",
        "Swap Cost",
        "Average Net SOC",
        "First Swap",
        "Last Swap",
    )
    book.close()

    csv_path = write_csv(rows, tmp_path / "out.csv")
    assert csv_path.read_text(encoding="utf-8").splitlines()[0].startswith("Period,Swap Station")
    assert "Matched Station Name" not in rows[0]


def test_combine_reports_totals_every_period_by_raw_station(tmp_path: Path) -> None:
    august = build_dump(tmp_path / "August_Swap_Data.xlsx", sample_rows())
    september = build_dump(tmp_path / "September_Swap_Data.xlsx", sample_rows())
    reports = [
        aggregate_swap_dump(august, year=2026),
        aggregate_swap_dump(september, year=2026),
    ]
    rows = combine_reports(reports)
    ahero = next(row for row in rows if row["Swap Station"] == "Ahero")
    assert ahero["Period"] == "2026-08, 2026-09"
    assert ahero["Net SOC"] == pytest.approx(1009.0)
    assert ahero["Billable SOC"] == pytest.approx(1008.0)
    assert ahero["Swap Cost"] == pytest.approx(1149.5)
    assert ahero["Swaps"] == 2
    assert ahero["Average Net SOC"] == pytest.approx(1009 / 2)
    assert {row["Swap Station"] for row in rows} == {"Ahero", "ahero", "Alendu"}


def test_discover_dumps_expands_dirs_and_globs(tmp_path: Path) -> None:
    dumps = tmp_path / "Swap Dump"
    dumps.mkdir()
    build_dump(dumps / "August_Swap_Data.xlsx", sample_rows())
    (dumps / "notes.txt").write_text("ignore me", encoding="utf-8")

    assert [item.name for item in discover_dumps([str(dumps)])] == ["August_Swap_Data.xlsx"]
    assert [item.name for item in discover_dumps([str(dumps / "*.xlsx")])] == [
        "August_Swap_Data.xlsx"
    ]
    assert discover_dumps([str(tmp_path / "missing.xlsx")]) == []
